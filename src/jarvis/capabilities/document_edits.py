"""Semantic document creation/editing for governed JARVIS Hands H2."""

from __future__ import annotations

import os
import pathlib
import tempfile
import time
from typing import Any

from jarvis.authority.types import ActionAttributes
from jarvis.capabilities.execution import PreparedCapability
from jarvis.capabilities.local_writes import (
    ApprovedWriteRootPolicy,
    LocalWriteValidationError,
    _contains_secret,
)
from jarvis.capabilities.models import (
    CapabilityDescriptor,
    CapabilityKind,
    CapabilityRequest,
    CapabilityResult,
    CapabilityStatus,
)

_MAX_TEXT = 20_000
_MAX_CELL_TEXT = 10_000


class DocumentEditValidationError(ValueError):
    pass


def _bounded_text(value: object, *, limit: int = _MAX_TEXT) -> str:
    text = str(value or "")
    if not text.strip() or len(text) > limit:
        raise DocumentEditValidationError("document text is empty or exceeds the bounded limit")
    if _contains_secret(text):
        raise DocumentEditValidationError("credential-like content is blocked from document edits")
    return text


def _atomic_save(target: pathlib.Path, saver) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=".jarvis-doc-", suffix=target.suffix, dir=target.parent)
    os.close(fd)
    temp = pathlib.Path(name)
    try:
        saver(temp)
        if not temp.is_file() or temp.stat().st_size == 0:
            raise OSError("document library produced an empty output")
        os.replace(temp, target)
    finally:
        if temp.exists():
            temp.unlink()


class DocumentEditExecutor:
    capability_key = "local:documents.edit"
    operations = (
        "add_pptx_text_slide",
        "append_docx_paragraph",
        "create_docx",
        "create_pptx",
        "create_xlsx",
        "set_xlsx_cell",
    )

    def __init__(self, roots: ApprovedWriteRootPolicy | None = None) -> None:
        self.roots = roots or ApprovedWriteRootPolicy()
        self.descriptor = CapabilityDescriptor.create(
            capability_id="documents.edit",
            source_id="local",
            kind=CapabilityKind.NATIVE_API,
            name="Semantic local document editing",
            description=(
                "Create and update DOCX, XLSX and PPTX files in governed write roots "
                "through mature document libraries rather than UI automation."
            ),
            operations=list(self.operations),
            metadata={"root_aliases": list(self.roots.aliases)},
            execution_enabled=True,
        )

    def prepare(self, request: CapabilityRequest) -> PreparedCapability:
        operation = request.operation
        if operation not in self.operations:
            raise DocumentEditValidationError("unsupported document edit operation")
        params = dict(request.parameters)
        root_alias = str(params.get("root") or "").strip().casefold()
        path = str(params.get("path") or "").strip()
        _, target, relative = self.roots.resolve(root_alias, path)
        extension = target.suffix.casefold()
        expected_ext = ".docx" if "docx" in operation else ".xlsx" if "xlsx" in operation else ".pptx"
        if extension != expected_ext:
            raise DocumentEditValidationError(
                f"{operation} requires a {expected_ext} target"
            )
        payload: dict[str, Any] = {"target": str(target)}
        normalized: dict[str, Any] = {"root": root_alias, "path": relative.as_posix()}

        if operation in {"create_docx", "append_docx_paragraph"}:
            text = _bounded_text(params.get("text"))
            normalized["text"] = text
            payload["text"] = text
        elif operation in {"create_pptx", "add_pptx_text_slide"}:
            title = _bounded_text(params.get("title"), limit=500)
            body = str(params.get("body") or "")
            if len(body) > _MAX_TEXT or _contains_secret(body):
                raise DocumentEditValidationError("slide body exceeds limit or looks credential-like")
            normalized.update(title=title, body=body)
            payload.update(title=title, body=body)
        elif operation == "set_xlsx_cell":
            sheet = str(params.get("sheet") or "Sheet").strip()
            cell = str(params.get("cell") or "").strip().upper()
            value = params.get("value")
            if not sheet or len(sheet) > 31:
                raise DocumentEditValidationError("worksheet name is invalid")
            import re

            if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,5}", cell):
                raise DocumentEditValidationError("cell must be a bounded A1-style address")
            if isinstance(value, str):
                if len(value) > _MAX_CELL_TEXT or _contains_secret(value):
                    raise DocumentEditValidationError("cell text exceeds limit or looks credential-like")
            elif value is not None and not isinstance(value, (int, float, bool)):
                raise DocumentEditValidationError("cell value type is not supported")
            normalized.update(sheet=sheet, cell=cell, value=value)
            payload.update(sheet=sheet, cell=cell, value=value)
        elif operation == "create_xlsx":
            sheet = str(params.get("sheet") or "Sheet").strip()
            if not sheet or len(sheet) > 31:
                raise DocumentEditValidationError("worksheet name is invalid")
            normalized["sheet"] = sheet
            payload["sheet"] = sheet

        summary = f"{operation.replace('_', ' ')} {root_alias}:{relative.as_posix()}"
        return PreparedCapability(
            request=request,
            target={"domain": "documents.edit", "root": root_alias, "path": relative.as_posix()},
            parameters=normalized,
            material_summary=summary,
            attributes=ActionAttributes(persistent_write=True),
            execution_payload=payload,
        )

    def execute(self, prepared: PreparedCapability) -> CapabilityResult:
        started = time.monotonic()
        try:
            data = self._execute(prepared.request.operation, prepared.execution_payload)
        except (DocumentEditValidationError, LocalWriteValidationError, OSError, ImportError) as exc:
            return CapabilityResult(
                status=CapabilityStatus.FAILED,
                capability_key=self.capability_key,
                operation=prepared.request.operation,
                data={},
                reason=str(exc),
                elapsed_ms=(time.monotonic() - started) * 1000.0,
            )
        return CapabilityResult(
            status=CapabilityStatus.SUCCEEDED,
            capability_key=self.capability_key,
            operation=prepared.request.operation,
            data={**data, "verification_passed": True},
            elapsed_ms=(time.monotonic() - started) * 1000.0,
            provenance=("python-docx/openpyxl/python-pptx semantic document libraries",),
        )

    @staticmethod
    def _execute(operation: str, payload: dict[str, Any]) -> dict[str, Any]:
        target = pathlib.Path(str(payload["target"]))
        if operation in {"create_docx", "append_docx_paragraph"}:
            from docx import Document

            if operation == "create_docx":
                if target.exists():
                    raise DocumentEditValidationError("DOCX target already exists")
                document = Document()
            else:
                if not target.is_file():
                    raise DocumentEditValidationError("DOCX target does not exist")
                document = Document(target)
            document.add_paragraph(str(payload["text"]))
            _atomic_save(target, document.save)
            check = Document(target)
            if not check.paragraphs or check.paragraphs[-1].text != str(payload["text"]):
                raise OSError("DOCX post-write verification failed")
            return {"path_exists": True, "paragraphs": len(check.paragraphs)}

        if operation in {"create_xlsx", "set_xlsx_cell"}:
            from openpyxl import Workbook, load_workbook

            if operation == "create_xlsx":
                if target.exists():
                    raise DocumentEditValidationError("XLSX target already exists")
                workbook = Workbook()
                workbook.active.title = str(payload["sheet"])
            else:
                if not target.is_file():
                    raise DocumentEditValidationError("XLSX target does not exist")
                workbook = load_workbook(target)
                sheet = str(payload["sheet"])
                if sheet not in workbook.sheetnames:
                    workbook.create_sheet(sheet)
                workbook[sheet][str(payload["cell"])] = payload["value"]
            _atomic_save(target, workbook.save)
            check = load_workbook(target, data_only=False)
            if operation == "set_xlsx_cell":
                actual = check[str(payload["sheet"])][str(payload["cell"])].value
                if actual != payload["value"]:
                    raise OSError("XLSX post-write verification failed")
            return {"path_exists": True, "sheets": tuple(check.sheetnames)}

        if operation in {"create_pptx", "add_pptx_text_slide"}:
            from pptx import Presentation

            if operation == "create_pptx":
                if target.exists():
                    raise DocumentEditValidationError("PPTX target already exists")
                presentation = Presentation()
            else:
                if not target.is_file():
                    raise DocumentEditValidationError("PPTX target does not exist")
                presentation = Presentation(target)
            layout = presentation.slide_layouts[1]
            slide = presentation.slides.add_slide(layout)
            slide.shapes.title.text = str(payload["title"])
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = str(payload["body"])
            _atomic_save(target, presentation.save)
            check = Presentation(target)
            if not check.slides or check.slides[-1].shapes.title.text != str(payload["title"]):
                raise OSError("PPTX post-write verification failed")
            return {"path_exists": True, "slides": len(check.slides)}

        raise DocumentEditValidationError("unsupported prepared document operation")
